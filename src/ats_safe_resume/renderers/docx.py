"""DOCX renderer using python-docx."""

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from ats_safe_resume.models import Resume
from ats_safe_resume.renderers.base import BaseRenderer
from ats_safe_resume import inline_md


class DocxRenderer(BaseRenderer):
    """Render resume to DOCX using python-docx."""

    def render(self, resume: Resume, output_path: Path) -> None:
        doc = Document()

        # Set default font
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Source Sans 3'
        font.size = Pt(10)

        # Name (centered, 18pt, bold)
        if resume.name:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(resume.name)
            run.bold = True
            run.font.size = Pt(18)
            run.font.name = 'Source Sans 3'

        # Title line (centered, 11pt, semibold)
        if resume.title_line:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(resume.title_line)
            run.bold = True
            run.font.size = Pt(11)
            run.font.name = 'Source Sans 3'

        # Contact (centered, 9pt)
        if resume.contact:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(resume.contact.to_ats_string())
            run.font.size = Pt(9)
            run.font.name = 'Source Sans 3'

        # Sections
        if resume.executive_profile:
            self._add_section_heading(doc, "Executive Profile")
            self._add_markdown_paragraph(doc, resume.executive_profile)

        if resume.core_expertise:
            self._add_section_heading(doc, "Core Expertise")
            self._add_markdown_paragraph(doc, resume.core_expertise)

        if resume.companies:
            self._add_section_heading(doc, "Professional Experience")
            for company in resume.companies:
                for position in company.positions:
                    loc = f" — {company.location}" if company.location else ""
                    p = doc.add_paragraph()
                    run = p.add_run(f"{company.name}{loc}")
                    run.bold = True
                    run.font.size = Pt(10.5)

                    title_parts = [position.title]
                    if position.subtitle:
                        title_parts.append(position.subtitle)
                    if position.start_date:
                        title_parts.append(f"({position.start_date} – {position.end_date or 'Present'})")

                    p = doc.add_paragraph()
                    run = p.add_run(" · ".join(title_parts))
                    run.bold = True
                    run.font.size = Pt(10)

                    if position.summary:
                        self._add_markdown_paragraph(doc, position.summary, Pt(10))

                    for bullet in position.bullets:
                        p = doc.add_paragraph(style='List Bullet')
                        self._add_markdown_runs(p, bullet)

        if resume.technical_skills:
            self._add_section_heading(doc, "Technical Skills")
            for skill in resume.technical_skills:
                p = doc.add_paragraph()
                cat_run = p.add_run(f"{skill.category}: ")
                cat_run.bold = True
                p.add_run(skill.skills)

        if resume.education:
            self._add_section_heading(doc, "Education")
            for edu in resume.education:
                p = doc.add_paragraph()
                inst_run = p.add_run(edu.institution)
                inst_run.bold = True
                parts = []
                if edu.degree:
                    parts.append(edu.degree)
                if edu.details:
                    parts.append(edu.details)
                if parts:
                    p.add_run(" — " + ", ".join(parts))

        if resume.crafted_footer:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("crafted with ats_safe_resume")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

        doc.save(str(output_path))

    def _add_section_heading(self, doc: Document, text: str) -> None:
        """Add a section heading with accent bar bottom border."""
        p = doc.add_paragraph()
        p.style = doc.styles['Heading 1']
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = 'Source Sans 3'
        # Accent bar (bottom border)
        pPr = p._element.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '555555')
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _add_markdown_paragraph(self, doc: Document, text: str, size: Pt | None = None) -> None:
        """Add a paragraph with inline markdown conversion."""
        p = doc.add_paragraph()
        self._add_markdown_runs(p, text, size)

    def _add_markdown_runs(self, paragraph, text: str, size: Pt | None = None) -> None:
        """Add runs with **bold**, *italic*, and links handling."""
        from docx.shared import RGBColor
        for token_text, is_bold, is_italic, url in inline_md.to_docx_runs(text):
            run = paragraph.add_run(token_text)
            run.bold = is_bold
            run.italic = is_italic
            run.font.name = 'Source Sans 3'
            if size:
                run.font.size = size
            if url:
                run.font.underline = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
